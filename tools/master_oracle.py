"""Open a .pptx produced by ruby_pptx and report the slide masters python-pptx sees.

python-pptx cannot *create* a slide master, so there is no differential oracle
for one. It reads them perfectly well, though, which is what this reports: the
masters, their layouts, the placeholders of each, and the theme colours and
fonts the master resolves to. A slide built on a generated layout is reported
too, so placeholder inheritance is visible.
"""

from __future__ import annotations

import json
import sys

import pptx
from pptx.oxml.ns import qn


def placeholders(collection):
    out = []
    for ph in collection:
        fmt = ph.placeholder_format
        out.append(
            {
                "type": str(fmt.type).split(".")[-1].split(" ")[0] if fmt.type else None,
                "idx": fmt.idx,
                "name": ph.name,
                "left": ph.left,
                "top": ph.top,
                "width": ph.width,
                "height": ph.height,
            }
        )
    return out


def theme_of(master_part):
    """The colour and font schemes of the theme part related to a master."""
    for rel in master_part.rels.values():
        if rel.reltype.endswith("/theme"):
            root = rel.target_part._element if hasattr(rel.target_part, "_element") else None
            if root is None:
                from lxml import etree

                root = etree.fromstring(rel.target_part.blob)
            elements = root.find(qn("a:themeElements"))
            colors = {}
            scheme = elements.find(qn("a:clrScheme"))
            for slot in scheme:
                tag = slot.tag.split("}")[-1]
                srgb = slot.find(qn("a:srgbClr"))
                sys_clr = slot.find(qn("a:sysClr"))
                if srgb is not None:
                    colors[tag] = srgb.get("val")
                elif sys_clr is not None:
                    colors[tag] = "sys:" + sys_clr.get("val")
            fonts = elements.find(qn("a:fontScheme"))
            return {
                "theme_name": root.get("name"),
                "scheme_name": scheme.get("name"),
                "colors": colors,
                "major": fonts.find(qn("a:majorFont")).find(qn("a:latin")).get("typeface"),
                "minor": fonts.find(qn("a:minorFont")).find(qn("a:latin")).get("typeface"),
            }
    return None


def main() -> int:
    prs = pptx.Presentation(sys.argv[1])
    masters = []
    for master in prs.slide_masters:
        masters.append(
            {
                "partname": str(master.part.partname),
                "placeholders": placeholders(master.placeholders),
                "theme": theme_of(master.part),
                "layouts": [
                    {
                        "name": layout.name,
                        "type": layout.element.get("type"),
                        "preserve": layout.element.get("preserve"),
                        "placeholders": placeholders(layout.placeholders),
                    }
                    for layout in master.slide_layouts
                ],
            }
        )

    slides = [
        {
            "layout": slide.slide_layout.name,
            "master": str(slide.slide_layout.slide_master.part.partname),
            "placeholders": placeholders(slide.placeholders),
            "texts": [
                ph.text_frame.text for ph in slide.placeholders if ph.has_text_frame
            ],
        }
        for slide in prs.slides
    ]

    json.dump({"masters": masters, "slides": slides}, sys.stdout)
    return 0


if __name__ == "__main__":
    sys.exit(main())
