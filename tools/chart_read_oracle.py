"""Read chart properties with python-pptx, for comparison with ruby_pptx.

The package differentials check what each library *writes*. They cannot see
what a library *reads*: an absent element whose schema default is true can be
reported as false and nothing in the file changes. This evaluates a list of
property expressions against one chart and prints the values as JSON, so a
spec can evaluate the same properties in Ruby and compare.

    python3 chart_read_oracle.py deck.pptx SLIDE_INDEX '["chart.chart_type", ...]'

Each expression is evaluated with `chart` bound. Enum members become their
names, category labels their text; an exception becomes "error:<class>".
"""

from __future__ import annotations

import enum
import json
import sys

import pptx


def normalise(value):
    if isinstance(value, enum.Enum):
        return value.name
    if isinstance(value, str):
        return str(value)
    if isinstance(value, (list, tuple)):
        return [normalise(v) for v in value]
    return value


def main() -> int:
    path, slide_index, expressions = sys.argv[1], int(sys.argv[2]), json.loads(sys.argv[3])
    prs = pptx.Presentation(path)
    chart = next(s for s in prs.slides[slide_index].shapes if s.has_chart).chart
    results = {}
    for expression in expressions:
        try:
            results[expression] = normalise(eval(expression, {"chart": chart}))  # noqa: S307
        except Exception as exc:  # noqa: BLE001 -- the error is the answer
            results[expression] = "error:" + type(exc).__name__
    json.dump(results, sys.stdout)
    return 0


if __name__ == "__main__":
    sys.exit(main())
