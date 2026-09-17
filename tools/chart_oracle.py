"""Open a .pptx produced by ruby_pptx and report what python-pptx sees in it.

The embedded chart workbook cannot be compared byte-for-byte, because
python-pptx writes it with XlsxWriter and ruby_pptx writes it directly. This
reports the semantics instead: the chart type, its categories and series, and
the cell grid of the workbook as a real spreadsheet reader sees it.
"""

from __future__ import annotations

import io
import json
import sys

import openpyxl
import pptx


def main() -> int:
    path = sys.argv[1]
    prs = pptx.Presentation(path)
    charts = []

    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_chart:
                continue
            chart = shape.chart
            xlsx_part = chart.part.chart_workbook.xlsx_part
            workbook = openpyxl.load_workbook(io.BytesIO(xlsx_part.blob))
            sheet = workbook[workbook.sheetnames[0]]
            charts.append(
                {
                    "chart_type": str(chart.chart_type),
                    "categories": [str(c) for c in chart.plots[0].categories],
                    "series": [
                        {"name": s.name, "values": list(s.values)} for s in chart.series
                    ],
                    "sheet_name": workbook.sheetnames[0],
                    "grid": [list(row) for row in sheet.iter_rows(values_only=True)],
                }
            )

    json.dump({"charts": charts}, sys.stdout)
    return 0


if __name__ == "__main__":
    sys.exit(main())
