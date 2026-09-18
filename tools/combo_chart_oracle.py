"""Report what python-pptx sees in a combo chart produced by ruby_pptx.

python-pptx can read a chart with several plots but never writes one, so there
is nothing to diff against. This reports the structure instead: the plots, the
series in each, and the axes.
"""

from __future__ import annotations

import json
import sys

import pptx


def main() -> int:
    prs = pptx.Presentation(sys.argv[1])
    charts = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_chart:
                continue
            chart = shape.chart
            charts.append(
                {
                    "plots": [
                        {
                            "kind": type(plot).__name__,
                            "series": [s.name for s in plot.series],
                            "categories": [str(c) for c in plot.categories],
                        }
                        for plot in chart.plots
                    ],
                    "value_axis_readable": _value_axis_readable(chart),
                }
            )
    json.dump({"charts": charts}, sys.stdout)
    return 0


def _value_axis_readable(chart) -> bool:
    try:
        chart.value_axis  # noqa: B018
        return True
    except Exception:  # noqa: BLE001
        return False


if __name__ == "__main__":
    sys.exit(main())
